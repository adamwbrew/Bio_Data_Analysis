"""Save figures, their supporting data, and reproducibility metadata."""

import importlib.metadata
import json
from pathlib import Path
import platform

import numpy as np
import pandas as pd
from openpyxl import Workbook


def save_figure(fig, name, output_dir, dpi=180):
    """Export a standalone PNG and a PDF with readable vector labels."""
    folder = Path(output_dir) / "figures"
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for suffix in ["png", "pdf"]:
        path = folder / f"{name}.{suffix}"
        fig.savefig(path, dpi=dpi, bbox_inches="tight", facecolor="white")
        paths.append(path)
    return paths


def _excel_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def export_tables(tables, output_dir, sources=None):
    """Write the actual chart data to gzip CSVs and a streaming Excel workbook.

    Large sheets are split before Excel's row limit. The workbook index records
    complete CSV paths; the input objects remain in the original HALO exports.
    """
    root = Path(output_dir)
    folder = root / "tables"
    folder.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    index = book.create_sheet("Index")
    index.append(["Table", "Rows", "Columns", "Compressed CSV", "Description"])
    used = {"Index", "Sources"}
    all_tables = dict(tables)
    if sources is not None:
        all_tables["Sources"] = sources
    for name, frame in all_tables.items():
        frame.to_csv(folder / f"{name}.csv.gz", index=False, compression="gzip")
        index.append([name, len(frame), len(frame.columns), f"tables/{name}.csv.gz",
                      "Exact values supporting the figures/tables; see notebook for units and definitions."])
        for start in range(0, max(1, len(frame)), 1_000_000):
            base = name[:25]
            suffix = "" if start == 0 else f"_{start // 1_000_000 + 1}"
            candidate = base + suffix
            if name == "Sources":
                candidate = "Sources"
            elif candidate in used:
                candidate = base[:22] + f"_{len(used)}"
            used.add(candidate)
            sheet = book.create_sheet(candidate)
            sheet.append(list(frame.columns))
            for row in frame.iloc[start:start + 1_000_000].itertuples(index=False, name=None):
                sheet.append([_excel_value(value) for value in row])
    destination = root / "visualization_data.xlsx"
    book.save(destination)
    return destination


def write_run_metadata(output_dir, parameters, manifest):
    """Record library versions, inputs and settings without machine user paths."""
    versions = {}
    for name in ["numpy", "pandas", "scipy", "matplotlib", "nbformat", "nbclient",
                 "nbconvert", "jupyterlab", "ipywidgets", "openpyxl"]:
        try:
            versions[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            versions[name] = "not installed"
    metadata = {"python": platform.python_version(), "packages": versions,
                "parameters": parameters, "input_manifest": manifest.to_dict(orient="records"),
                "interpretation": "Descriptive object-level analysis; biological replicate identities unconfirmed.",
                "coordinates": "Export units; x/y are bounding-box midpoints. No physical scale inferred.",
                "source_images_available": False}
    destination = Path(output_dir) / "run_metadata.json"
    destination.write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n")
    return destination
